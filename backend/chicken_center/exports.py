"""
Excel exports for Chicken Center reports. Reads the accounting core services
and streams .xlsx using the getvalue() pattern (BytesIO bytes, not the buffer
object) that the Feeds exports proved reliable.
"""
from io import BytesIO
from datetime import date as date_cls
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from django.http import HttpResponse

from accounting.services import pnl_detailed, cash_book
from .models import Sale
from .views import compute_ageing

BIZ = 'chicken_center'
HEAD_FILL = PatternFill(start_color='1F6F43', end_color='1F6F43', fill_type='solid')
HEAD_FONT = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)


def _xlsx_response(wb, filename):
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    resp = HttpResponse(
        out.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center')


def _dates(request):
    df = request.query_params.get('from') or date_cls.today().replace(day=1).isoformat()
    dt = request.query_params.get('to') or date_cls.today().isoformat()
    return df, dt


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_pnl(request):
    df, dt = _dates(request)
    data = pnl_detailed(df, dt, business=BIZ)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Profit & Loss'
    ws.append([f'Profit & Loss  {df} to {dt}'])
    ws['A1'].font = BOLD
    ws.append([])
    ws.append(['Income', 'Amount'])
    _style_header(ws, ws.max_row)
    for r in data['income_lines']:
        ws.append([r['name'], float(r['amount'])])
    ws.append(['Total Income', float(data['income'])])
    ws[ws.max_row][0].font = BOLD
    ws.append([])
    ws.append(['Expenses', 'Amount'])
    _style_header(ws, ws.max_row)
    for r in data['expense_lines']:
        ws.append([r['name'], float(r['amount'])])
    ws.append(['Total Expenses', float(data['expense'])])
    ws[ws.max_row][0].font = BOLD
    ws.append([])
    ws.append(['Net Profit', float(data['profit'])])
    ws[ws.max_row][0].font = BOLD
    ws[ws.max_row][1].font = BOLD
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 16
    return _xlsx_response(wb, f'chicken_pnl_{df}_to_{dt}.xlsx')


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_cash_book(request):
    df, dt = _dates(request)
    data = cash_book(df, dt, business=BIZ)
    wb = openpyxl.Workbook()
    first = True
    for code, cb in data.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = cb['account'][:31]
        first = False
        ws.append([f"{cb['account']}  {df} to {dt}"])
        ws['A1'].font = BOLD
        ws.append(['Opening', float(cb['opening'])])
        ws.append(['Received', float(cb['receipts'])])
        ws.append(['Paid', float(cb['payments'])])
        ws.append(['Closing', float(cb['closing'])])
        ws.append([])
        ws.append(['Date', 'Type', 'Details', 'In', 'Out', 'Balance'])
        _style_header(ws, ws.max_row)
        for row in cb['rows']:
            ws.append([str(row['date']), row['type'],
                       row['party'] or row['narration'],
                       float(row['in']), float(row['out']), float(row['balance'])])
        for col, w in zip('ABCDEF', (12, 12, 30, 12, 12, 14)):
            ws.column_dimensions[col].width = w
    return _xlsx_response(wb, f'chicken_cashbook_{df}_to_{dt}.xlsx')


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_pending_bills(request):
    qs = (Sale.objects
          .filter(status__in=['PENDING', 'PARTLY'])
          .exclude(voucher__is_reversed=True)
          .select_related('party')
          .order_by('date', 'id'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Pending Bills'
    ws.append(['Date', 'Bill #', 'Customer', 'Amount', 'Received', 'Due', 'Status'])
    _style_header(ws)
    total = Decimal('0')
    for s in qs:
        ws.append([str(s.date), s.id, s.party.name, float(s.amount),
                   float(s.amount_received), float(s.amount_due), s.status])
        total += s.amount_due
    ws.append([])
    ws.append(['', '', '', '', '', float(total), 'TOTAL DUE'])
    ws[ws.max_row][5].font = BOLD
    for col, w in zip('ABCDEFG', (12, 8, 24, 14, 14, 14, 12)):
        ws.column_dimensions[col].width = w
    return _xlsx_response(wb, 'chicken_pending_bills.xlsx')


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_ageing(request):
    data = compute_ageing()  # same bucketing as the on-screen report
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ageing'
    ws.append(['Customer', '0-7 days', '8-15 days', '16-30 days', '31+ days', 'Total', 'Oldest (days)'])
    _style_header(ws)
    for r in data['rows']:
        ws.append([r['party_name'], float(r['0-7']), float(r['8-15']),
                   float(r['16-30']), float(r['31+']), float(r['total']), r['oldest_days']])
    t = data['totals']
    ws.append(['TOTAL', float(t['0-7']), float(t['8-15']), float(t['16-30']),
               float(t['31+']), float(data['grand_total']), ''])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for col, w in zip('ABCDEFG', (24, 12, 12, 12, 12, 14, 14)):
        ws.column_dimensions[col].width = w
    return _xlsx_response(wb, 'chicken_ageing.xlsx')
